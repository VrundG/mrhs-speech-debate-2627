"""Create the privacy-minimized member dataset used by the dashboard."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl


HISTORY_LABELS = {
    "speech and debate-n. mecklenburg": "N. Mecklenburg Viking Classic",
    "speech & debate-laird lewis": "Laird Lewis Invitational",
    "speech & debate-corona rostrensis": "Corona Rostrensis",
    "speech & debate-ashville tournament": "Cougar Classic at Asheville",
    "speech & debate-toc online 12/6-8": "TOC Online",
    "speech & debate-world school tutorial 12/20": "World Schools Tutorial",
    "speech and debate-durham": "Cavalier Invitational at Durham Academy",
    "marvin ridge hs-tfl-speech & debate-2026": "TFL State Championship",
    "marvin ridge hs-cuthbertson classic-speech & debate": "Cuthbertson Classic",
    "marvin ridge hs-speech and debate carolina district": "Carolina West Districts",
    "marvin ridge hs-columbia online 2026-speech & debate": "Columbia University Online Invitational",
    "marvin ridge hs-stanford online 2026-speech & debate - pf & vld": "Stanford Online Invitational",
    "marvin ridge hs-stanford online 2026-speech & debate - sppech & nld": "Stanford Online Invitational",
}


def clean_name(value: object) -> str:
    text = re.sub(r"\s+", " ", str(value or "")).strip()
    if "," in text:
        last, first = [part.strip() for part in text.split(",", 1)]
        text = f"{first} {last}"
    return " ".join(part.title() for part in text.split())


def iso_date(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        try:
            return datetime.fromisoformat(value).date().isoformat()
        except ValueError:
            return None
    return None


def build(source: Path, destination: Path) -> None:
    workbook = openpyxl.load_workbook(source, data_only=True, read_only=True)
    membership_sheet = workbook["members total"]
    header = [str(cell.value or "").strip().lower() for cell in next(membership_sheet.iter_rows())]
    indexes = {name: header.index(name) for name in ("student id", "student name")}

    members_by_key: dict[str, dict[str, object]] = {}
    for row in membership_sheet.iter_rows(values_only=True):
        raw_name = row[indexes["student name"]] if len(row) > indexes["student name"] else None
        if not raw_name or str(raw_name).strip().lower() == "student name":
            continue
        name = clean_name(raw_name)
        student_id = str(row[indexes["student id"]] or "").replace(".0", "")
        key = student_id or name.casefold()
        members_by_key.setdefault(
            key,
            {
                "id": hashlib.sha256(key.encode()).hexdigest()[:12],
                "name": name,
                "role": "Student",
                "memberType": "Returning member",
                "membershipFee": 45,
                "membershipStatus": "Unpaid",
                "eventHistory": [],
                "tournamentHistory": [],
            },
        )

    name_to_member = {str(member["name"]).casefold(): member for member in members_by_key.values()}
    history: dict[str, dict[str, dict[str, object]]] = defaultdict(dict)

    for sheet in workbook.worksheets:
        rows = list(sheet.iter_rows(values_only=True))
        if not rows:
            continue
        headers = [str(value or "").strip().lower() for value in rows[0]]
        student_indexes = [index for index, value in enumerate(headers) if value in {"student", "student name"}]
        item_indexes = [index for index, value in enumerate(headers) if value in {"item", "item name"}]
        date_indexes = [index for index, value in enumerate(headers) if value in {"date", "date/time"}]
        if not student_indexes or not item_indexes:
            continue

        for row in rows[1:]:
            for student_index in student_indexes:
                if student_index >= len(row) or not row[student_index]:
                    continue
                member = name_to_member.get(clean_name(row[student_index]).casefold())
                if not member:
                    continue
                for item_index in item_indexes:
                    if item_index >= len(row) or not row[item_index]:
                        continue
                    raw_item = str(row[item_index]).strip().lower()
                    tournament = HISTORY_LABELS.get(raw_item)
                    if not tournament:
                        continue
                    observed_date = None
                    for date_index in date_indexes:
                        if date_index < len(row):
                            observed_date = iso_date(row[date_index])
                            if observed_date:
                                break
                    history[str(member["id"])][tournament] = {
                        "tournament": tournament,
                        "date": observed_date,
                    }

    members = []
    for member in members_by_key.values():
        member["tournamentHistory"] = sorted(
            history[str(member["id"])].values(),
            key=lambda entry: str(entry["date"] or ""),
            reverse=True,
        )
        members.append(member)

    members.sort(key=lambda member: str(member["name"]).casefold())
    if len(members) != 158:
        raise ValueError(f"Expected 158 deduplicated members, found {len(members)}")

    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(members, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {len(members)} deduplicated members to {destination}")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    args = parser.parse_args()
    build(args.source, args.destination)


if __name__ == "__main__":
    main()
